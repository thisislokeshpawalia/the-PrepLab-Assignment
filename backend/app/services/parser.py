import io
from fastapi import HTTPException
import pymupdf
from docx import Document


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extracts text cleanly from PDF bytes using PyMuPDF."""
    try:
        doc = pymupdf.open(stream=file_bytes, filetype="pdf")
        text_parts = []
        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text("text")
            if text:
                text_parts.append(text.strip())
        doc.close()
        full_text = "\n\n".join(text_parts).strip()
        if not full_text:
            raise HTTPException(
                status_code=400,
                detail="Could not extract text from the PDF. The file may be image-only (scanned) or empty."
            )
        return full_text
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to parse PDF file: {str(e)}")


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extracts text cleanly from DOCX bytes using python-docx."""
    try:
        file_stream = io.BytesIO(file_bytes)
        doc = Document(file_stream)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        
        # Also extract table text if any
        table_text = []
        for table in doc.tables:
            for row in table.rows:
                row_str = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_str:
                    table_text.append(row_str)
                    
        full_text = "\n\n".join(paragraphs + table_text).strip()
        if not full_text:
            raise HTTPException(
                status_code=400,
                detail="Could not extract text from DOCX file. File appears to be empty."
            )
        return full_text
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to parse DOCX file: {str(e)}")


def parse_resume(filename: str, file_bytes: bytes) -> str:
    """Dispatches to appropriate parser based on file extension."""
    lower_name = filename.lower()
    if lower_name.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif lower_name.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    elif lower_name.endswith(".txt"):
        return file_bytes.decode("utf-8", errors="ignore").strip()
    else:
        raise HTTPException(
            status_code=400,
            detail="Unsupported file format. Please upload a PDF (.pdf) or Word document (.docx)."
        )
