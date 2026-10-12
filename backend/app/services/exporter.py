import io
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from app.schemas import ResumeSections


def export_resume_to_docx(resume: ResumeSections) -> io.BytesIO:
    """Generates a professionally formatted Word Document."""
    doc = Document()
    
    # Header: Name & Contact
    name = resume.name or "Candidate Name"
    doc.add_heading(name, level=0)
    
    if resume.contact_info:
        p_contact = doc.add_paragraph(resume.contact_info)
        p_contact.runs[0].font.size = 10
    
    # Summary
    if resume.summary:
        doc.add_heading("Professional Summary", level=1)
        doc.add_paragraph(resume.summary)
        
    # Skills
    if resume.skills:
        doc.add_heading("Technical Skills & Core Competencies", level=1)
        doc.add_paragraph(" • ".join(resume.skills))
        
    # Experience
    if resume.experience:
        doc.add_heading("Professional Experience", level=1)
        for exp in resume.experience:
            header_text = f"{exp.role} — {exp.company}"
            if exp.duration:
                header_text += f" ({exp.duration})"
            doc.add_heading(header_text, level=2)
            if exp.location:
                p_loc = doc.add_paragraph(exp.location)
                p_loc.italic = True
            for bullet in exp.bullets:
                doc.add_paragraph(bullet, style='List Bullet')
                
    # Projects
    if resume.projects:
        doc.add_heading("Key Projects", level=1)
        for proj in resume.projects:
            proj_title = proj.name
            if proj.technologies:
                proj_title += f" [{proj.technologies}]"
            doc.add_heading(proj_title, level=2)
            for bullet in proj.bullets:
                doc.add_paragraph(bullet, style='List Bullet')
                
    # Education
    if resume.education:
        doc.add_heading("Education", level=1)
        for edu in resume.education:
            edu_title = f"{edu.degree} — {edu.institution}"
            if edu.year:
                edu_title += f" ({edu.year})"
            doc.add_heading(edu_title, level=2)
            if edu.details:
                doc.add_paragraph(edu.details)
                
    # Certifications
    if resume.certifications:
        doc.add_heading("Certifications", level=1)
        for cert in resume.certifications:
            doc.add_paragraph(cert, style='List Bullet')

    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream


def export_resume_to_pdf(resume: ResumeSections) -> io.BytesIO:
    """Generates an ATS-friendly, clean PDF using ReportLab."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40
    )
    
    styles = getSampleStyleSheet()
    
    # Custom Styles
    name_style = ParagraphStyle(
        'DocName',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#0f172a'),
        alignment=1  # Centered
    )
    contact_style = ParagraphStyle(
        'DocContact',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#475569'),
        alignment=1
    )
    section_heading = ParagraphStyle(
        'DocSection',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#1e40af'),
        spaceBefore=10,
        spaceAfter=4,
        keepWithNext=True
    )
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13,
        textColor=colors.HexColor('#1e293b')
    )
    bullet_style = ParagraphStyle(
        'DocBullet',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12.5,
        leftIndent=15,
        textColor=colors.HexColor('#334155'),
        spaceAfter=3
    )
    subhead_style = ParagraphStyle(
        'DocSubhead',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#0f172a'),
        spaceBefore=4
    )

    story = []
    
    # Name & Contact
    story.append(Paragraph(resume.name or "Candidate Name", name_style))
    if resume.contact_info:
        story.append(Paragraph(resume.contact_info, contact_style))
    story.append(Spacer(1, 10))
    
    # Summary
    if resume.summary:
        story.append(Paragraph("PROFESSIONAL SUMMARY", section_heading))
        story.append(Paragraph(resume.summary, body_style))
        story.append(Spacer(1, 6))
        
    # Skills
    if resume.skills:
        story.append(Paragraph("SKILLS & COMPETENCIES", section_heading))
        story.append(Paragraph(" • ".join(resume.skills), body_style))
        story.append(Spacer(1, 6))
        
    # Experience
    if resume.experience:
        story.append(Paragraph("EXPERIENCE", section_heading))
        for exp in resume.experience:
            header_str = f"<b>{exp.role}</b> — {exp.company}"
            if exp.duration:
                header_str += f" | {exp.duration}"
            story.append(Paragraph(header_str, subhead_style))
            for b in exp.bullets:
                story.append(Paragraph(f"• {b}", bullet_style))
        story.append(Spacer(1, 6))
        
    # Projects
    if resume.projects:
        story.append(Paragraph("PROJECTS", section_heading))
        for proj in resume.projects:
            p_str = f"<b>{proj.name}</b>"
            if proj.technologies:
                p_str += f" <i>({proj.technologies})</i>"
            story.append(Paragraph(p_str, subhead_style))
            for b in proj.bullets:
                story.append(Paragraph(f"• {b}", bullet_style))
        story.append(Spacer(1, 6))

    # Education
    if resume.education:
        story.append(Paragraph("EDUCATION", section_heading))
        for edu in resume.education:
            e_str = f"<b>{edu.degree}</b> — {edu.institution}"
            if edu.year:
                e_str += f" ({edu.year})"
            story.append(Paragraph(e_str, subhead_style))
            if edu.details:
                story.append(Paragraph(edu.details, body_style))
        story.append(Spacer(1, 6))

    doc.build(story)
    buffer.seek(0)
    return buffer
